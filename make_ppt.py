from pptx import Presentation
from pptx.util import Inches, Pt

# Create a new presentation
prs = Presentation()

# --- SLIDE 1: The Hook ---
slide_layout = prs.slide_layouts[0] # Title Slide layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Dimensional Analysis Semantic Checker"
subtitle.text = "Preventing the next Mars Climate Orbiter disaster\n\nTeam: Balraj, Ayansh, Shreyansh"

# --- SLIDE 2: Current Architecture ---
slide_layout = prs.slide_layouts[1] # Title and Content layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Current Architecture (The 'Before' State)"
tf = content.text_frame
tf.text = "The Frontend Pipeline (Completed)"

p = tf.add_paragraph()
p.text = "Lexer & Parser: Successfully reads .dim source files."
p.level = 1

p = tf.add_paragraph()
p.text = "Abstract Syntax Tree (AST): Correctly built from syntax."
p.level = 1

p = tf.add_paragraph()
p.text = "Symbol Table: Populated with variables and basic types."
p.level = 1

p = tf.add_paragraph()
p.text = "Current Status: The compiler perfectly understands grammar, but not physics."
p.level = 0
p.font.bold = True

# --- SLIDE 3: The Gap ---
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Gap: Syntax vs. Semantics"
tf = content.text_frame
tf.text = "Test Harness Results:"

p = tf.add_paragraph()
p.text = "5 Valid Tests: PASSED (Syntax is correct)"
p.level = 1

p = tf.add_paragraph()
p.text = "4 Invalid Tests: FAILED (Semantics are missing)"
p.level = 1

p = tf.add_paragraph()
p.text = "Example: 'let bad = distance + time' is accepted."
p.level = 1

p = tf.add_paragraph()
p.text = "The frontend parses it without complaining because the syntax is valid. But physically, you cannot add meters to seconds!"
p.level = 0
p.font.bold = True

# --- SLIDE 4: The Transition ---
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Solution: Closing the Gap"
tf = content.text_frame
tf.text = "Next Steps in the Pipeline:"

p = tf.add_paragraph()
p.text = "Semantic Checker (Ayansh): Walks the AST, calculates dimensional vectors, and throws compile-time errors on unit mismatches."
p.level = 1

p = tf.add_paragraph()
p.text = "Code Generator (Shreyansh): Takes the validated AST, strips away unit metadata, and emits pure, zero-overhead C code."
p.level = 1

p = tf.add_paragraph()
p.text = "Result: A robust compiler that catches physical impossibilities before the code ever runs."
p.level = 0
p.font.bold = True

# Save the presentation
prs.save('Presentation_Before.pptx')
print("Success! Presentation saved as 'Presentation_Before.pptx'")
